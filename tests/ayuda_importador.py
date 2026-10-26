"""
tests/ayuda_importador.py
-------------------------
Ayudas para las pruebas del importador: crear archivos Excel / CSV de ejemplo
y analizarlos con pocas líneas.
"""

import csv
import os

import openpyxl

from app.services.importador import lector
from app.services.importador.analizador import analizar
from app.services.importador.catalogo import cargar_catalogo
from app.services.importador.mapeo import adivinar_mapeo, detectar_fila_encabezado
from app.services.importador.modelos import ConfigImportacion


class Literal(str):
    """Texto que se guarda tal cual, aunque empiece con "=" (si no, openpyxl lo guarda como fórmula)."""


def crear_xlsx(ruta: str, hojas: dict[str, list[list]], ocultas=()) -> str:
    """hojas = {"nombre de hoja": [[fila 1], [fila 2], ...]}"""
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(fila)
            for celda, valor in zip(hoja[hoja.max_row], fila):
                if isinstance(valor, Literal):
                    celda.data_type = "s"
        if nombre in ocultas:
            hoja.sheet_state = "hidden"
    libro.save(ruta)
    return ruta


def crear_csv(ruta: str, filas: list[list], separador: str = ";", codificacion: str = "cp1252") -> str:
    with open(ruta, "w", encoding=codificacion, newline="") as f:
        csv.writer(f, delimiter=separador).writerows(filas)
    return ruta


def analizar_hojas(carpeta: str, hojas: dict[str, list[list]], mapeo=None, **config):
    """
    Escribe el Excel, adivina las columnas igual que el asistente y lo analiza.
    `config` son los campos de ConfigImportacion (modo_deposito, crear_depositos, ...).
    """
    ruta = crear_xlsx(os.path.join(carpeta, "inventario.xlsx"), hojas)
    if mapeo is None:
        primera = next(iter(hojas))
        previa = lector.leer_previa(ruta, primera)
        fila = detectar_fila_encabezado(previa) or 1
        mapeo = adivinar_mapeo(previa[fila - 1])
    cfg = ConfigImportacion(ruta=ruta, hojas=list(hojas), mapeo=mapeo, **config)
    return analizar(cfg, cargar_catalogo())

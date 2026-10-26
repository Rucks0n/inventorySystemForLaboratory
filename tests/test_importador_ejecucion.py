"""
tests/test_importador_ejecucion.py
----------------------------------
Pruebas de lo que SÍ guarda en la base: que los datos queden bien, que repetir la
importación no duplique, y sobre todo que ante un fallo o una cancelación la base
quede EXACTAMENTE como estaba.
"""

import os
import sqlite3
import unittest

import openpyxl

from app.controllers.inventario_controller import DepositosController
from app.services.importador import ejecutor, perfiles
from app.services.importador.analizador import analizar
from app.services.importador.catalogo import cargar_catalogo
from app.services.importador.mapeo import adivinar_mapeo, mapeo_de_perfil
from app.services.importador.modelos import ConfigImportacion, ErrorImportacion, ImportacionCancelada
from app.services.importador.plantilla import crear_plantilla
from app.services.importador.reporte_errores import escribir_reporte
from config.constants import IMP_DEPOSITO_UNICO, IMP_DUP_ACTUALIZAR, IMP_DUP_OMITIR
from db.migraciones import aplicar_pendientes
from tests.ayuda_importador import Literal, analizar_hojas, crear_csv
from tests.base_prueba import BaseConBD

CABECERA = ["Nombre", "Cantidad", "Referencia", "Ubicación", "Tipo", "Unidad", "Observaciones"]


class BaseEjecucion(BaseConBD):
    def setUp(self):
        super().setUp()
        aplicar_pendientes()
        for nombre in ("Depósito1", "Laboratorio A"):
            DepositosController().agregar({"nombre": nombre})
        self.carpeta = self._carpeta.name
        self.tamano_lote_original = ejecutor.TAMANO_LOTE

    def tearDown(self):
        ejecutor.TAMANO_LOTE = self.tamano_lote_original
        super().tearDown()

    # ── atajos ─────────────────────────────────────────────────────────────
    def analizar(self, hojas, **config):
        return analizar_hojas(self.carpeta, hojas, **config)

    def tipo_id(self, nombre):
        return next(t.id for t in cargar_catalogo().tipos_por_id.values() if t.nombre == nombre)

    def deposito_id(self, nombre):
        return self.consulta("SELECT id FROM depositos WHERE nombre = ?", (nombre,))[0]["id"]

    def items(self):
        return [dict(f) for f in self.consulta("""
            SELECT d.nombre AS deposito, i.nombre, i.referencia, i.cantidad_total, i.cantidad_disp,
                   i.cantidad_prest, i.ubicacion, i.observaciones, t.nombre AS tipo, i.unidad
            FROM items i JOIN depositos d ON d.id = i.deposito_id
            LEFT JOIN tipos_item t ON t.id = i.tipo_id ORDER BY i.id""")]

    def crear_item(self, deposito="Depósito1", nombre="Sensor", referencia=None, total=5, prestados=0,
                   ubicacion=None, observaciones=None, tipo="Herramienta"):
        return self.ejecutar(
            "INSERT INTO items (deposito_id, nombre, referencia, cantidad_total, cantidad_disp, cantidad_prest,"
            " ubicacion, observaciones, tipo_id) VALUES (?,?,?,?,?,?,?,?,?)",
            (self.deposito_id(deposito), nombre, referencia, total, total - prestados, prestados,
             ubicacion, observaciones, self.tipo_id(tipo)))

    def importar(self, analisis, politica=IMP_DUP_OMITIR, **kw):
        return ejecutor.ejecutar(analisis, politica, **kw)


class PruebasGuardado(BaseEjecucion):

    def test_guarda_todos_los_datos_y_deja_todo_disponible(self):
        a = self.analizar({"Depósito1": [CABECERA,
                                         ["Micrómetro", 3, "M-1", "Estante 1", "Herramienta", None, "Calibrado"],
                                         ["Aceite", 500, None, "Estante 2", "Consumible", "mL", None]]})
        r = self.importar(a)
        self.assertEqual((r.creados, r.actualizados, r.omitidos, r.conflictos), (2, 0, 0, []))
        self.assertEqual(self.items(), [
            {"deposito": "Depósito1", "nombre": "Micrómetro", "referencia": "M-1", "cantidad_total": 3,
             "cantidad_disp": 3, "cantidad_prest": 0, "ubicacion": "Estante 1", "observaciones": "Calibrado",
             "tipo": "Herramienta", "unidad": None},
            {"deposito": "Depósito1", "nombre": "Aceite", "referencia": None, "cantidad_total": 500,
             "cantidad_disp": 500, "cantidad_prest": 0, "ubicacion": "Estante 2", "observaciones": None,
             "tipo": "Consumible", "unidad": "mL"}])

    def test_el_tipo_por_defecto_solo_se_aplica_a_filas_sin_tipo(self):
        a = self.analizar({"Depósito1": [["Nombre", "Cantidad", "Tipo"],
                                         ["Pesa", 1, None], ["Grafito", 9, "Consumible"]]},
                          tipo_defecto_id=self.tipo_id("Herramienta"))
        self.importar(a)
        self.assertEqual([(i["nombre"], i["tipo"]) for i in self.items()],
                         [("Pesa", "Herramienta"), ("Grafito", "Consumible")])

    def test_sin_tipo_por_defecto_los_items_quedan_sin_tipo(self):
        self.importar(self.analizar({"Depósito1": [["Nombre", "Cantidad"], ["Pesa", 1]]}))
        self.assertIsNone(self.items()[0]["tipo"])

    def test_reparte_los_items_entre_depositos(self):
        a = self.analizar({"Depósito1": [["Nombre", "Cantidad"], ["A", 1]],
                           "Laboratorio A": [["Nombre", "Cantidad"], ["B", 2]]})
        self.importar(a)
        self.assertEqual([(i["deposito"], i["nombre"]) for i in self.items()],
                         [("Depósito1", "A"), ("Laboratorio A", "B")])

    def test_crea_los_depositos_que_faltan_si_se_permitio(self):
        a = self.analizar({"Bodega Nueva": [["Nombre", "Cantidad"], ["A", 1]]}, crear_depositos=True)
        r = self.importar(a)
        self.assertEqual(r.depositos_creados, ["Bodega Nueva"])
        self.assertEqual(self.items()[0]["deposito"], "Bodega Nueva")

    def test_importa_desde_csv(self):
        ruta = crear_csv(os.path.join(self.carpeta, "lab.csv"),
                         [["Nombre", "Cantidad", "Ubicación"], ["Ácido nítrico", 2, "Estante ñ"]])
        cfg = ConfigImportacion(ruta=ruta, hojas=["lab"], mapeo={"nombre": "nombre", "cantidad": "cantidad",
                                                                  "ubicacion": "ubicacion"},
                                modo_deposito=IMP_DEPOSITO_UNICO, deposito_unico="Depósito1")
        self.importar(analizar(cfg, cargar_catalogo()))
        self.assertEqual([(i["nombre"], i["cantidad_total"], i["ubicacion"]) for i in self.items()],
                         [("Ácido nítrico", 2, "Estante ñ")])

    def test_sin_filas_validas_no_hace_nada_ni_crea_respaldo(self):
        a = self.analizar({"Depósito1": [["Nombre", "Cantidad"], ["Roto", "abc"]]})
        with self.assertRaises(ErrorImportacion):
            self.importar(a)
        self.assertFalse(os.path.exists(self.carpeta_respaldos))

    def test_un_inventario_grande(self):
        filas = [["Nombre", "Cantidad", "Referencia"]] + [[f"Ítem {i}", i % 50, f"R{i}"] for i in range(1, 5001)]
        r = self.importar(self.analizar({"Depósito1": filas}))
        self.assertEqual(r.creados, 5000)
        self.assertEqual(self.consulta("SELECT COUNT(*) AS n FROM items")[0]["n"], 5000)


class PruebasRepetirImportacion(BaseEjecucion):

    def test_subir_dos_veces_el_mismo_archivo_no_duplica(self):
        hojas = {"Depósito1": [["Nombre", "Cantidad", "Referencia"], ["Cable USB", 5, "A"], ["Pesa", 2, None]]}
        self.importar(self.analizar(hojas))
        r = self.importar(self.analizar(hojas))
        self.assertEqual((r.creados, r.omitidos), (0, 2))
        self.assertEqual(len(self.items()), 2)

    def test_omitir_no_modifica_lo_existente(self):
        self.crear_item(nombre="Sensor", total=5, ubicacion="Caja 1")
        r = self.importar(self.analizar({"Depósito1": [["Nombre", "Cantidad", "Ubicación"], ["sensor", 99, "Otra"]]}))
        self.assertEqual(r.omitidos, 1)
        self.assertEqual((self.items()[0]["cantidad_total"], self.items()[0]["ubicacion"]), (5, "Caja 1"))

    def test_corregir_el_excel_y_volver_a_subir_solo_agrega_lo_que_faltaba(self):
        primera = {"Depósito1": [["Nombre", "Cantidad"], ["Bueno", 1], ["Con error", "abc"]]}
        self.importar(self.analizar(primera))                               # "Con error" se rechaza
        corregida = {"Depósito1": [["Nombre", "Cantidad"], ["Bueno", 1], ["Con error", 4]]}
        r = self.importar(self.analizar(corregida))
        self.assertEqual((r.creados, r.omitidos), (1, 1))
        self.assertEqual([i["nombre"] for i in self.items()], ["Bueno", "Con error"])

    def test_se_crea_una_copia_de_seguridad_con_los_datos_de_antes(self):
        self.crear_item(nombre="Previo")
        r = self.importar(self.analizar({"Depósito1": [["Nombre", "Cantidad"], ["Nuevo", 1]]}))
        self.assertTrue(os.path.exists(r.respaldo))
        self.assertIn("antes_de_importar", r.respaldo)
        copia = sqlite3.connect(r.respaldo)
        self.assertEqual([f[0] for f in copia.execute("SELECT nombre FROM items")], ["Previo"])
        copia.close()
        self.assertEqual(len(self.items()), 2)


class PruebasActualizar(BaseEjecucion):

    def test_actualiza_cantidad_conservando_lo_prestado(self):
        self.crear_item(nombre="Pesa", total=5, prestados=2)                # 3 disponibles, 2 prestadas
        r = self.importar(self.analizar({"Depósito1": [["Nombre", "Cantidad"], ["Pesa", 10]]}), IMP_DUP_ACTUALIZAR)
        self.assertEqual((r.actualizados, r.conflictos), (1, []))
        i = self.items()[0]
        self.assertEqual((i["cantidad_total"], i["cantidad_prest"], i["cantidad_disp"]), (10, 2, 8))

    def test_una_celda_vacia_no_borra_lo_que_ya_estaba(self):
        self.crear_item(nombre="Pesa", total=5, ubicacion="Estante 1", observaciones="calibrada")
        self.importar(self.analizar({"Depósito1": [["Nombre", "Cantidad", "Ubicación", "Observaciones"],
                                                   ["Pesa", 6, None, None]]}), IMP_DUP_ACTUALIZAR)
        i = self.items()[0]
        self.assertEqual((i["cantidad_total"], i["ubicacion"], i["observaciones"]), (6, "Estante 1", "calibrada"))

    def test_actualiza_los_campos_que_si_vienen(self):
        self.crear_item(nombre="Pesa", total=5, ubicacion="Estante 1")
        self.importar(self.analizar({"Depósito1": [["Nombre", "Cantidad", "Ubicación", "Tipo", "Unidad"],
                                                   ["Pesa", 5, "Estante 9", "Consumible", "g"]]}), IMP_DUP_ACTUALIZAR)
        i = self.items()[0]
        self.assertEqual((i["ubicacion"], i["tipo"], i["unidad"]), ("Estante 9", "Consumible", "g"))

    def test_si_el_archivo_es_igual_a_lo_guardado_cuenta_como_omitido(self):
        self.crear_item(nombre="Pesa", total=5)
        r = self.importar(self.analizar({"Depósito1": [["Nombre", "Cantidad"], ["Pesa", 5]]}), IMP_DUP_ACTUALIZAR)
        self.assertEqual((r.actualizados, r.omitidos), (0, 1))

    def test_no_baja_la_cantidad_por_debajo_de_lo_prestado(self):
        self.crear_item(nombre="Pesa", total=5, prestados=4)
        r = self.importar(self.analizar({"Depósito1": [["Nombre", "Cantidad"], ["Pesa", 2], ["Nueva", 1]]}),
                          IMP_DUP_ACTUALIZAR)
        self.assertEqual((r.creados, r.actualizados, len(r.conflictos)), (1, 0, 1))
        self.assertIn("menor que la que está prestada", r.conflictos[0].mensaje)
        self.assertEqual(self.items()[0]["cantidad_total"], 5)              # no se tocó

    def test_no_cambia_a_un_tipo_de_otro_comportamiento_con_unidades_prestadas(self):
        self.crear_item(nombre="Pesa", total=5, prestados=1, tipo="Herramienta")
        r = self.importar(self.analizar({"Depósito1": [["Nombre", "Cantidad", "Tipo"], ["Pesa", 5, "Consumible"]]}),
                          IMP_DUP_ACTUALIZAR)
        self.assertIn("mientras haya unidades prestadas", r.conflictos[0].mensaje)
        self.assertEqual(self.items()[0]["tipo"], "Herramienta")

    def test_si_hay_dos_iguales_en_la_base_no_adivina_cual_actualizar(self):
        self.crear_item(nombre="Pesa", total=1)
        self.crear_item(nombre="Pesa", total=2)
        r = self.importar(self.analizar({"Depósito1": [["Nombre", "Cantidad"], ["Pesa", 9]]}), IMP_DUP_ACTUALIZAR)
        self.assertIn("2 ítems iguales", r.conflictos[0].mensaje)
        self.assertEqual([i["cantidad_total"] for i in self.items()], [1, 2])


class PruebasIntegridad(BaseEjecucion):
    """Si algo sale mal, la base queda como estaba: nada a medias."""

    def estado(self):
        return (self.items(), [dict(f) for f in self.consulta("SELECT nombre FROM depositos ORDER BY id")])

    def test_un_fallo_a_mitad_deshace_todo_incluidos_los_depositos_nuevos(self):
        ejecutor.TAMANO_LOTE = 2
        a = self.analizar({"Bodega Nueva": [["Nombre", "Cantidad"], ["A", 1], ["B", 2], ["C", 3], ["D", 4], ["E", 5]]},
                          crear_depositos=True)
        a.filas[3].tipo_id = 9999                     # un tipo inexistente: falla en el SEGUNDO lote
        antes = self.estado()
        with self.assertRaises(ErrorImportacion) as ctx:
            self.importar(a)
        self.assertIn("Tu inventario quedó como estaba", str(ctx.exception))
        self.assertEqual(self.estado(), antes)        # ni ítems del primer lote, ni la bodega nueva

    def test_cancelar_a_mitad_deshace_todo(self):
        ejecutor.TAMANO_LOTE = 2
        self.crear_item(nombre="Previo")
        a = self.analizar({"Depósito1": [["Nombre", "Cantidad"]] + [[f"X{i}", i] for i in range(1, 6)]})
        antes = self.estado()
        llamadas = []

        def cancelar_en_el_segundo_lote():
            llamadas.append(1)
            return len(llamadas) >= 2                 # el primer lote ya se escribió cuando llega aquí

        with self.assertRaises(ImportacionCancelada):
            self.importar(a, cancelado=cancelar_en_el_segundo_lote)
        self.assertEqual(self.estado(), antes)

    def test_informa_el_progreso_del_guardado(self):
        ejecutor.TAMANO_LOTE = 2
        a = self.analizar({"Depósito1": [["Nombre", "Cantidad"]] + [[f"X{i}", i] for i in range(1, 6)]})
        avances = []
        self.importar(a, progreso=lambda msg, hecho, total: avances.append((hecho, total)))
        self.assertEqual(avances, [(2, 5), (4, 5), (5, 5)])

    def test_la_importacion_no_rompe_las_relaciones_de_la_base(self):
        self.importar(self.analizar({"Depósito1": [["Nombre", "Cantidad", "Tipo"], ["A", 1, "Consumible"]]}))
        self.assertEqual(self.consulta("PRAGMA foreign_key_check"), [])


class PruebasVistaPrevia(BaseEjecucion):

    def test_simular_cuenta_lo_que_pasaria_sin_guardar_nada(self):
        self.crear_item(nombre="Existente", total=5)
        a = self.analizar({"Depósito1": [["Nombre", "Cantidad"], ["Existente", 5], ["Nuevo", 1]],
                           "Bodega Nueva": [["Nombre", "Cantidad"], ["Otro", 2]]}, crear_depositos=True)
        antes_items, antes_depositos = len(self.items()), self.consulta("SELECT COUNT(*) AS n FROM depositos")[0]["n"]

        plan = ejecutor.simular(a, IMP_DUP_OMITIR)

        self.assertEqual((len(plan.crear), len(plan.omitir), len(plan.actualizar)), (2, 1, 0))
        self.assertEqual(plan.depositos_nuevos, ["Bodega Nueva"])
        self.assertEqual(len(self.items()), antes_items)
        self.assertEqual(self.consulta("SELECT COUNT(*) AS n FROM depositos")[0]["n"], antes_depositos)

    def test_simular_con_actualizar_muestra_actualizaciones_y_conflictos(self):
        self.crear_item(nombre="A", total=5)
        self.crear_item(nombre="B", total=5, prestados=5)
        a = self.analizar({"Depósito1": [["Nombre", "Cantidad"], ["A", 7], ["B", 1]]})
        plan = ejecutor.simular(a, IMP_DUP_ACTUALIZAR)
        self.assertEqual((len(plan.actualizar), len(plan.conflictos)), (1, 1))


class PruebasPerfiles(BaseEjecucion):

    def test_guardar_listar_y_reemplazar_sin_importar_mayusculas(self):
        perfiles.guardar("Laboratorio de Física", {"nombre": "insumo", "cantidad": "n"})
        perfiles.guardar("laboratorio de física", {"nombre": "insumo", "cantidad": "n", "unidad": "um"})
        lista = perfiles.listar()
        self.assertEqual(len(lista), 1)
        self.assertEqual(lista[0]["mapeo"]["unidad"], "um")

    def test_no_guarda_perfiles_vacios(self):
        for nombre, mapeo in (("", {"nombre": "x"}), ("Lab", {})):
            with self.assertRaises(ValueError):
                perfiles.guardar(nombre, mapeo)

    def test_un_perfil_guardado_se_aplica_solo_al_subir_un_archivo_con_esas_columnas(self):
        perfiles.guardar("Física", {"nombre": "insumo", "cantidad": "n"})
        nombre, mapeo = mapeo_de_perfil(["N", "Insumo", "Notas"], perfiles.listar())
        self.assertEqual((nombre, mapeo), ("Física", {"nombre": "insumo", "cantidad": "n"}))
        # Con ese mapeo el archivo se importa aunque las columnas tengan nombres raros
        a = self.analizar({"Depósito1": [["N", "Insumo"], [4, "Termómetro"]]}, mapeo=mapeo)
        self.assertEqual([(f.nombre, f.cantidad) for f in a.filas], [("Termómetro", 4)])


class PruebasReporteYPlantilla(BaseEjecucion):

    def test_el_reporte_lista_los_problemas_con_su_fila_y_su_contenido(self):
        a = self.analizar({"Depósito1": [["Nombre", "Cantidad"], ["Bueno", 1], ["Malo", "abc"],
                                         [Literal("=CMD()"), "-"]]})
        ruta = os.path.join(self.carpeta, "problemas.xlsx")
        n = escribir_reporte(ruta, a, conflictos=[])
        self.assertEqual(n, 2)
        filas = list(openpyxl.load_workbook(ruta)["Problemas"].iter_rows(values_only=True))
        self.assertEqual(filas[0], ("Hoja", "Fila", "Gravedad", "Motivo", "Contenido de la fila"))
        self.assertEqual([(f[0], f[1], f[2]) for f in filas[1:]], [("Depósito1", 3, "Error"), ("Depósito1", 4, "Error")])
        self.assertEqual(filas[1][4], "Malo | abc")
        self.assertTrue(filas[2][4].startswith("'=CMD()"))        # nunca se escribe como fórmula

    def test_el_reporte_incluye_conflictos_y_problemas_de_hoja_completa(self):
        self.crear_item(nombre="Pesa", total=5, prestados=5)
        a = self.analizar({"Depósito1": [["Nombre", "Cantidad"], ["Pesa", 1]], "Fantasma": [["Nombre", "Cantidad"], ["X", 1]]})
        plan = ejecutor.simular(a, IMP_DUP_ACTUALIZAR)
        ruta = os.path.join(self.carpeta, "problemas.xlsx")
        escribir_reporte(ruta, a, plan.conflictos)
        filas = list(openpyxl.load_workbook(ruta)["Problemas"].iter_rows(values_only=True))[1:]
        self.assertEqual([f[1] for f in filas], [2, "(toda la hoja)"])

    def test_la_plantilla_se_reconoce_sola_y_trae_reglas(self):
        ruta = os.path.join(self.carpeta, "plantilla.xlsx")
        crear_plantilla(ruta, ["Herramienta", "Consumible"])
        libro = openpyxl.load_workbook(ruta)
        self.assertEqual(libro.sheetnames, ["Inventario", "Instrucciones"])
        encabezados = [c.value for c in libro["Inventario"][1]]
        self.assertEqual(adivinar_mapeo(encabezados), {
            "nombre": "nombre", "cantidad": "cantidad", "referencia": "referencia", "ubicacion": "ubicacion",
            "tipo": "tipo", "unidad": "unidad", "observaciones": "observaciones"})
        reglas = {v.type: v for v in libro["Inventario"].data_validations.dataValidation}
        self.assertEqual(set(reglas), {"whole", "list"})
        self.assertIn("Herramienta,Consumible", reglas["list"].formula1)

    def test_flujo_completo_llenando_la_plantilla(self):
        ruta = os.path.join(self.carpeta, "plantilla.xlsx")
        crear_plantilla(ruta, ["Herramienta", "Consumible"])
        libro = openpyxl.load_workbook(ruta)
        hoja = libro["Inventario"]
        hoja.append(["Micrómetro", 2, "M-1", "Estante 1", "Herramienta", None, None])
        hoja.append(["Aceite", 500, None, "Estante 2", "Consumible", "mL", "SAE 40"])
        libro.save(ruta)

        cfg = ConfigImportacion(ruta=ruta, hojas=["Inventario"], modo_deposito=IMP_DEPOSITO_UNICO,
                                deposito_unico="Laboratorio A",
                                mapeo=adivinar_mapeo([c.value for c in hoja[1]]))
        a = analizar(cfg, cargar_catalogo())
        self.assertEqual(a.incidencias, [])
        self.importar(a)
        self.assertEqual([(i["deposito"], i["nombre"], i["tipo"], i["unidad"]) for i in self.items()],
                         [("Laboratorio A", "Micrómetro", "Herramienta", None),
                          ("Laboratorio A", "Aceite", "Consumible", "mL")])


if __name__ == "__main__":
    unittest.main()

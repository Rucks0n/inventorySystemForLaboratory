"""
app/services/importador/reporte_errores.py
------------------------------------------
Genera un Excel con TODOS los problemas encontrados al revisar el archivo:
hoja, número de fila, gravedad, motivo y el contenido original de la fila.

Sirve para corregir el archivo y volver a subirlo: como los ítems que ya se
importaron se omiten, se puede subir el archivo completo otra vez sin duplicar.
"""

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from app.services.importador.modelos import Analisis, Incidencia
from config.constants import IMP_ERROR

ENCABEZADOS = ("Hoja", "Fila", "Gravedad", "Motivo", "Contenido de la fila")
ANCHOS      = (22, 8, 11, 90, 70)


def escribir_reporte(ruta_destino: str, analisis: Analisis,
                     conflictos: list[Incidencia] | tuple = ()) -> int:
    """Guarda el reporte en `ruta_destino` (.xlsx). Retorna cuántos problemas escribió."""
    orden_hojas = {h: i for i, h in enumerate(analisis.config.hojas)}
    problemas = sorted(
        list(analisis.incidencias) + list(conflictos),
        key=lambda p: (orden_hojas.get(p.hoja, 999), p.fila),
    )

    libro = Workbook(write_only=True)
    hoja = libro.create_sheet("Problemas")
    for i, ancho in enumerate(ANCHOS):
        hoja.column_dimensions["ABCDE"[i]].width = ancho
    hoja.freeze_panes = "A2"

    hoja.append([_celda(hoja, t, negrita=True) for t in ENCABEZADOS])
    for p in problemas:
        contenido = analisis.originales.get((p.hoja, p.fila))
        hoja.append([
            _celda(hoja, p.hoja),
            _celda(hoja, p.fila if p.fila else "(toda la hoja)"),
            _celda(hoja, "Error" if p.nivel == IMP_ERROR else "Aviso"),
            _celda(hoja, p.mensaje),
            _celda(hoja, " | ".join("" if v is None else str(v) for v in contenido) if contenido else ""),
        ])
    libro.save(ruta_destino)
    return len(problemas)


def _celda(hoja, valor, negrita: bool = False) -> WriteOnlyCell:
    # Un texto que empiece con "=" se interpretaría como fórmula: se escribe literal
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@"):
        valor = "'" + valor
    celda = WriteOnlyCell(hoja, value=valor)
    celda.alignment = Alignment(vertical="top", wrap_text=True)
    if negrita:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="2E5090")
    return celda

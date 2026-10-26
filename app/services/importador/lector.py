"""
app/services/importador/lector.py
---------------------------------
Lee archivos Excel (.xlsx / .xlsm) y CSV FILA POR FILA.

Para inventarios muy grandes lo importante es no cargar todo el archivo en
memoria: aquí cada fila se entrega y se olvida. Lo único que se guarda es la
fila que se está revisando.

Cosas que este lector resuelve (y que suelen dar problemas):
  - Archivos abiertos en Excel, dañados o con contraseña → mensaje claro.
  - .xls antiguos → se explica cómo convertirlos (no se pueden leer directo).
  - Excel con filas vacías al final (Excel a veces "recuerda" miles de filas
    en blanco) → se saltan.
  - Archivos cuya "dimensión" interna está mal escrita (pasa con archivos
    generados por otros programas) → se recalcula.
  - Fórmulas: se lee el ÚLTIMO VALOR calculado que Excel guardó.
  - CSV de Excel en español: separador ";" y codificación Windows (cp1252).
"""

import csv
import itertools
import os
import zipfile
from contextlib import ExitStack, contextmanager
from typing import Iterator, NamedTuple

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from app.services.importador.modelos import ErrorImportacion

EXTENSIONES_EXCEL = (".xlsx", ".xlsm")
EXTENSION_CSV     = ".csv"
FORMATOS_ACEPTADOS = "Excel (.xlsx) o CSV (.csv)"

_NO_SOPORTADOS = {
    ".xls":  "Los archivos .xls (Excel antiguo) no se pueden leer directamente. "
             "Ábrelo en Excel y guárdalo como .xlsx.",
    ".xlsb": "Los archivos .xlsb no se pueden leer. Ábrelo en Excel y guárdalo como .xlsx.",
    ".ods":  "Los archivos .ods no se pueden leer. Guárdalo como .xlsx o .csv desde LibreOffice.",
}


class HojaInfo(NamedTuple):
    nombre: str
    visible: bool


# ─────────────────────────────────────────────────────────────────────────────
#  API pública
# ─────────────────────────────────────────────────────────────────────────────

def validar_archivo(ruta: str) -> str:
    """Comprueba que el archivo existe y es de un tipo aceptado. Retorna su extensión."""
    if not os.path.isfile(ruta):
        raise ErrorImportacion(f"No encuentro el archivo:\n{ruta}")
    extension = os.path.splitext(ruta)[1].lower()
    if extension in _NO_SOPORTADOS:
        raise ErrorImportacion(_NO_SOPORTADOS[extension])
    if extension not in EXTENSIONES_EXCEL + (EXTENSION_CSV,):
        raise ErrorImportacion(f"Formato no aceptado ({extension or 'sin extensión'}). "
                               f"Usa un archivo {FORMATOS_ACEPTADOS}.")
    return extension


def fila_vacia(valores: list) -> bool:
    """True si todas las celdas están vacías o solo tienen espacios."""
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in valores)


class Archivo:
    """
    Un archivo abierto para leer varias hojas.

    ABRIR UN EXCEL GRANDE CUESTA (openpyxl carga de una vez todos los textos del
    libro, aunque solo se lea una hoja). Por eso el archivo se abre UNA sola vez y
    de ahí se lee todo lo que haga falta:

        with Archivo(ruta) as archivo:
            hojas  = archivo.hojas()
            previa = archivo.previa("Depósito1")
            for numero, celdas in archivo.filas("Depósito1", desde=2): ...

    Mientras está abierto, Excel puede mostrar el archivo como "en uso".
    """

    def __init__(self, ruta: str):
        self.ruta = ruta
        self.extension = validar_archivo(ruta)
        self._pila = ExitStack()
        self._libro = None

    def __enter__(self) -> "Archivo":
        if self.extension != EXTENSION_CSV:
            self._libro = self._pila.enter_context(_libro(self.ruta))
        return self

    def __exit__(self, *excepcion):
        self._pila.close()

    def hojas(self) -> list[HojaInfo]:
        """Hojas del archivo. Un CSV cuenta como una sola hoja con el nombre del archivo."""
        if self._libro is None:
            return [HojaInfo(os.path.splitext(os.path.basename(self.ruta))[0], True)]
        return [HojaInfo(ws.title, ws.sheet_state == "visible") for ws in self._libro.worksheets]

    def previa(self, hoja: str, max_filas: int = 40) -> list[list]:
        """Las primeras filas (incluidas las vacías): para detectar encabezados y mostrar ejemplos."""
        filas = self._filas(hoja, 1)
        try:
            return [valores for _, valores in itertools.islice(filas, max_filas)]
        finally:
            filas.close()

    def filas(self, hoja: str, desde: int) -> Iterator[tuple[int, list]]:
        """
        Entrega (número de fila, celdas) a partir de la fila `desde` (la primera es la 1).
        Las filas completamente vacías se saltan, pero el número de fila sigue siendo el
        real del Excel, para poder decirle al usuario "revisa la fila 87".
        """
        for numero, valores in self._filas(hoja, desde):
            if not fila_vacia(valores):
                yield numero, valores

    def _filas(self, hoja: str, desde: int) -> Iterator[tuple[int, list]]:
        if self._libro is None:
            yield from _filas_csv(self.ruta, desde)
        else:
            yield from _filas_excel(self._libro, hoja, desde)


# Atajos para leer una sola vez (abren y cierran el archivo en cada llamada)

def listar_hojas(ruta: str) -> list[HojaInfo]:
    with Archivo(ruta) as archivo:
        return archivo.hojas()


def leer_previa(ruta: str, hoja: str, max_filas: int = 40) -> list[list]:
    with Archivo(ruta) as archivo:
        return archivo.previa(hoja, max_filas)


def leer_filas(ruta: str, hoja: str, desde_fila: int) -> Iterator[tuple[int, list]]:
    with Archivo(ruta) as archivo:
        yield from archivo.filas(hoja, desde_fila)


# ─────────────────────────────────────────────────────────────────────────────
#  Internos
# ─────────────────────────────────────────────────────────────────────────────


@contextmanager
def _libro(ruta: str):
    try:
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    except PermissionError:
        raise ErrorImportacion("No pude abrir el archivo. ¿Está abierto en Excel? "
                               "Ciérralo e inténtalo de nuevo.") from None
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError):
        raise ErrorImportacion("El archivo no parece un Excel .xlsx válido. Puede estar dañado "
                               "o protegido con contraseña.") from None
    try:
        yield libro
    finally:
        libro.close()


def _filas_excel(libro, hoja: str, desde: int) -> Iterator[tuple[int, list]]:
    if hoja not in libro.sheetnames:
        raise ErrorImportacion(f"El archivo no tiene una hoja llamada «{hoja}».")
    ws = libro[hoja]
    ws.reset_dimensions()                  # no confiar en la dimensión que declara el archivo
    for numero, fila in enumerate(ws.iter_rows(min_row=desde, values_only=True), start=desde):
        yield numero, list(fila)


# ── CSV ──────────────────────────────────────────────────────────────────────

def _detectar_codificacion(ruta: str) -> str:
    """UTF-8 (con o sin marca BOM) si es válido; si no, la codificación de Excel en Windows."""
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            with open(ruta, encoding=codificacion) as f:
                while f.read(1 << 20):
                    pass
            return codificacion
        except UnicodeDecodeError:
            continue
    return "latin-1"                       # nunca falla


def _detectar_separador(primera_linea: str) -> str:
    """El Excel en español guarda CSV con ';'; otros programas usan ',' o tabulador."""
    candidatos = {s: primera_linea.count(s) for s in (";", ",", "\t", "|")}
    mejor = max(candidatos, key=candidatos.get)
    return mejor if candidatos[mejor] > 0 else ","


def _filas_csv(ruta: str, desde: int) -> Iterator[tuple[int, list]]:
    try:
        codificacion = _detectar_codificacion(ruta)
        with open(ruta, encoding=codificacion, newline="") as f:
            muestra = next((l for l in itertools.islice(f, 20) if l.strip()), "")
            f.seek(0)
            lector = csv.reader(f, delimiter=_detectar_separador(muestra))
            for numero, fila in enumerate(lector, start=1):
                if numero >= desde:
                    yield numero, [c if c != "" else None for c in fila]
    except PermissionError:
        raise ErrorImportacion("No pude abrir el archivo. ¿Está abierto en otro programa?") from None
    except csv.Error as e:
        raise ErrorImportacion(f"El archivo CSV no se pudo leer: {e}") from None

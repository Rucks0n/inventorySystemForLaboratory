"""
app/services/importador/plantilla.py
------------------------------------
Genera la plantilla de Excel que se puede descargar desde el asistente.

La plantilla es una AYUDA, no una obligación: el importador entiende archivos
con otras columnas y en otro orden. Pero si todos los laboratorios empiezan
desde la misma plantilla, casi no habrá errores.

Los nombres de las columnas están pensados para que el importador los reconozca
solos (ver campos.py). Si se cambian aquí, comprobar que sigan reconociéndose.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# (encabezado, obligatorio, ancho)
COLUMNAS = (
    ("Nombre",        True,  34),
    ("Cantidad",      True,  11),
    ("Referencia",    False, 18),
    ("Ubicación",     False, 18),
    ("Tipo",          False, 16),
    ("Unidad",        False, 12),
    ("Observaciones", False, 40),
)
FILAS_CON_REGLAS = 20000

INSTRUCCIONES = (
    "CÓMO LLENAR LA HOJA «Inventario»",
    "",
    "• Una fila por ítem. No dejes filas de títulos ni de totales entre los datos.",
    "• Nombre y Cantidad son obligatorias. Las demás columnas son opcionales:",
    "  si tu laboratorio no usa código o unidad, borra esa columna o déjala vacía.",
    "• Cantidad: solo números enteros, sin unidades ni decimales (escribe 500, no «500 mL»).",
    "• Referencia: el código, serial o modelo del ítem, si lo tiene.",
    "• Tipo: elige de la lista (Herramienta se presta y se devuelve; Consumible se gasta).",
    "• Unidad: para consumibles, la unidad en que cuentas (mL, g, unidades).",
    "  Si necesitas decimales, usa una unidad más pequeña (mL en vez de L).",
    "",
    "VARIOS DEPÓSITOS",
    "• Una hoja por depósito: el nombre de la hoja es el nombre del depósito, o",
    "• Una sola hoja con una columna «Depósito» que diga a cuál pertenece cada fila.",
    "",
    "El orden de las columnas no importa; el sistema las reconoce por su nombre.",
    "Puedes volver a subir el mismo archivo: lo que ya existe no se duplica.",
)


def crear_plantilla(ruta_destino: str, tipos: list[str]) -> None:
    """Guarda la plantilla en `ruta_destino` (.xlsx). `tipos` son los nombres de tipo válidos."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Inventario"

    for columna, (titulo, obligatorio, ancho) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=columna, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="2E5090" if obligatorio else "6B7FA8")
        celda.alignment = Alignment(horizontal="center", vertical="center")
        hoja.column_dimensions[celda.column_letter].width = ancho
    hoja.freeze_panes = "A2"
    hoja.row_dimensions[1].height = 24

    letra = {titulo: hoja.cell(row=1, column=i).column_letter
             for i, (titulo, _, _) in enumerate(COLUMNAS, start=1)}

    cantidades = DataValidation(
        type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
        showErrorMessage=True, errorTitle="Cantidad no válida",
        error="Escribe un número entero (0 o más), sin unidades ni decimales.")
    cantidades.add(f"{letra['Cantidad']}2:{letra['Cantidad']}{FILAS_CON_REGLAS}")
    hoja.add_data_validation(cantidades)

    lista = ",".join(tipos)
    if tipos and len(lista) <= 250:                        # Excel limita la lista a 255 caracteres
        lista_tipos = DataValidation(type="list", formula1=f'"{lista}"', allow_blank=True,
                                     showErrorMessage=True, errorTitle="Tipo no válido",
                                     error="Elige uno de los tipos de la lista.")
        lista_tipos.add(f"{letra['Tipo']}2:{letra['Tipo']}{FILAS_CON_REGLAS}")
        hoja.add_data_validation(lista_tipos)

    ayuda = libro.create_sheet("Instrucciones")
    ayuda.column_dimensions["A"].width = 95
    for fila, texto in enumerate(INSTRUCCIONES, start=1):
        celda = ayuda.cell(row=fila, column=1, value=texto)
        if texto.isupper() or texto.startswith("CÓMO") or texto == "VARIOS DEPÓSITOS":
            celda.font = Font(bold=True, color="2E5090")

    libro.save(ruta_destino)
